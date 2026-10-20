"""Build the static seed data from public sources.

Run once (and again when a source publishes a new version):

    python -m ingest.build_seed

Writes to data/seed/:
  ilo_scores.json   GenAI exposure for every ISCO-08 4-digit occupation (ILO WP140, CC BY 4.0)
  sdg_targets.json  Official SDG goal and target wording (UN Statistics Division SDG API)
  occupations.json  ESCO occupations used by the app, with ISCO code, essential skills
                    and the skill clusters those skills map to
"""
from __future__ import annotations

import io
import json
import time
from datetime import date
from pathlib import Path

import httpx

SEED = Path(__file__).resolve().parent.parent / "data" / "seed"
ILO_XLSX = "https://raw.githubusercontent.com/pgmyrek/2025_GenAI_scores_ISCO08/main/Final_Scores_ISCO08_Gmyrek_et_al_2025.xlsx"
ESCO_API = "https://ec.europa.eu/esco/api"
SDG_API = "https://unstats.un.org/SDGAPI/v1/sdg"

# Occupations people often hold today. They are offered as "current job" choices
# even when the live ESCO search is unavailable.
COMMON_CURRENT = [
    "office clerk", "bookkeeper", "secretary", "administrative assistant",
    "customer service representative", "call centre agent", "data entry clerk",
    "shop assistant", "receptionist", "warehouse worker",
]


def get_json(client: httpx.Client, url: str, params: dict | None = None, tries: int = 5):
    """GET with retries; ESCO sometimes drops connections under load."""
    for attempt in range(tries):
        try:
            r = client.get(url, params=params)
            r.raise_for_status()
            return r.json()
        except (httpx.HTTPError, ValueError):
            if attempt == tries - 1:
                raise
            time.sleep(2 * (attempt + 1))


def build_ilo(client: httpx.Client) -> dict:
    from openpyxl import load_workbook  # dev dependency, only needed here

    raw = client.get(ILO_XLSX, follow_redirects=True)
    raw.raise_for_status()
    ws = load_workbook(io.BytesIO(raw.content), read_only=True).active
    rows = ws.iter_rows(values_only=True)
    head = [str(h) for h in next(rows)]
    col = {name: head.index(name) for name in
           ("ISCO_08", "Title", "Task_ISCO", "score_2025", "mean_score_2025",
            "SD_2025", "mean_score_2023", "potential25")}
    occ: dict[str, dict] = {}
    for row in rows:
        if row[col["ISCO_08"]] is None:
            continue
        code = str(row[col["ISCO_08"]]).split(".")[0].zfill(4)
        o = occ.setdefault(code, {
            "title": row[col["Title"]],
            "mean_2025": round(float(row[col["mean_score_2025"]]), 3),
            "sd_2025": round(float(row[col["SD_2025"]]), 3),
            "mean_2023": round(float(row[col["mean_score_2023"]]), 3) if row[col["mean_score_2023"]] is not None else None,
            "category": row[col["potential25"]],
            "tasks": [],
        })
        o["tasks"].append({"task": str(row[col["Task_ISCO"]]).strip(), "score": round(float(row[col["score_2025"]]), 3)})
    for o in occ.values():
        o["tasks"].sort(key=lambda t: -t["score"])
    return {
        "_source": "Gmyrek et al. (2025), Generative AI and Jobs: A Refined Global Index of Occupational Exposure, ILO Working Paper 140. CC BY 4.0. " + ILO_XLSX,
        "_retrieved": date.today().isoformat(),
        "occupations": occ,
    }


def build_sdg(client: httpx.Client) -> dict:
    goals = get_json(client, f"{SDG_API}/Goal/List", {"includechildren": "false"})
    targets = get_json(client, f"{SDG_API}/Target/List", {"includechildren": "false"})
    return {
        "_source": "United Nations Statistics Division, SDG API " + SDG_API,
        "_retrieved": date.today().isoformat(),
        "goals": {g["code"]: g["title"] for g in goals},
        "targets": {t["code"]: {"goal": t["goal"], "title": t["title"]} for t in targets},
    }


def esco_occupation(client: httpx.Client, title: str) -> dict | None:
    res = get_json(client, f"{ESCO_API}/search",
                   {"text": title, "type": "occupation", "language": "en", "limit": 10})
    hits = res["_embedded"]["results"]
    hit = next((h for h in hits if h["title"].lower() == title.lower()), None)
    if not hit:
        return None
    o = get_json(client, f"{ESCO_API}/resource/occupation", {"uri": hit["uri"], "language": "en"})
    links = o.get("_links", {})
    desc = (o.get("description") or {}).get("en", {}).get("literal", "")
    return {
        "uri": hit["uri"],
        "title": hit["title"],
        "code": o.get("code"),
        "isco": (o.get("code") or "")[:4],
        "description": desc,
        "essential": [s["title"] for s in links.get("hasEssentialSkill", [])],
        "optional": [s["title"] for s in links.get("hasOptionalSkill", [])],
    }


def map_to_clusters(skill_titles: list[str], clusters: list[dict]) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for s in skill_titles:
        low = s.lower()
        for c in clusters:
            if any(k in low for k in c["esco_keywords"]):
                out.setdefault(c["id"], []).append(s)
    return out


def build_occupations(client: httpx.Client) -> dict:
    roles = json.loads((SEED / "roles.json").read_text(encoding="utf-8"))["roles"]
    samples = json.loads((SEED / "samples.json").read_text(encoding="utf-8"))
    clusters = json.loads((SEED / "clusters.json").read_text(encoding="utf-8"))["clusters"]
    titles = [r["esco_title"] for r in roles]
    titles += [s["esco_title"] for r in roles for s in r.get("ladder", [])]
    titles += [p["current_esco_title"] for p in samples["personas"] if p.get("current_esco_title")]
    titles += COMMON_CURRENT
    occ, missing = {}, []
    for t in dict.fromkeys(titles):
        o = esco_occupation(client, t)
        time.sleep(0.5)
        if not o:
            missing.append(t)
            continue
        o["clusters"] = map_to_clusters(o["essential"], clusters)
        occ[t] = o
        print(f"  ESCO {t!r}: ISCO {o['isco']}, {len(o['essential'])} essential skills")
    if missing:
        print("  Not found in ESCO (skipped):", ", ".join(missing))
    return {
        "_source": "ESCO classification v1.2, European Commission. " + ESCO_API,
        "_retrieved": date.today().isoformat(),
        "occupations": occ,
    }


def main() -> None:
    with httpx.Client(timeout=60, headers={"User-Agent": "workforce-futures/0.1 (hackathon prototype)"}) as client:
        print("ILO exposure scores…")
        ilo = build_ilo(client)
        (SEED / "ilo_scores.json").write_text(json.dumps(ilo, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"  {len(ilo['occupations'])} occupations")
        print("UN SDG targets…")
        sdg = build_sdg(client)
        (SEED / "sdg_targets.json").write_text(json.dumps(sdg, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"  {len(sdg['targets'])} targets")
        print("ESCO occupations…")
        occ = build_occupations(client)
        (SEED / "occupations.json").write_text(json.dumps(occ, ensure_ascii=False, indent=1), encoding="utf-8")


if __name__ == "__main__":
    main()
